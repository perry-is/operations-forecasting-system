"""Read-only intake and validation for the human-maintained source workbook."""

from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from .models import DataSet, Item, Order, PurchaseOrder, Shipment, Usage


def number(value) -> Decimal:
    result = Decimal(str(value))
    if not result.is_finite():
        raise ValueError("Quantities and costs must be finite")
    return result


def _as_date(value, field: str, optional: bool = False) -> date | None:
    if value in (None, "") and optional:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        converted = from_excel(value)
        return converted.date() if isinstance(converted, datetime) else converted
    try:
        return date.fromisoformat(str(value))
    except ValueError as exc:
        raise ValueError(f"{field} must be an ISO date") from exc


def _records(sheet) -> list[dict]:
    values = sheet.iter_rows(values_only=True)
    headers = next(values, None)
    if not headers or any(value in (None, "") for value in headers):
        raise ValueError(f"{sheet.title} needs a complete header row")
    names = [str(value).strip() for value in headers]
    if len(names) != len(set(names)):
        raise ValueError(f"{sheet.title} contains duplicate column names")
    result = []
    for row_number, row in enumerate(values, start=2):
        if all(value is None for value in row):
            continue
        if len(row) != len(names):
            raise ValueError(f"{sheet.title} row {row_number} does not match its headers")
        result.append(dict(zip(names, row)))
    return result


def _required(record: dict, fields: set[str], sheet: str) -> None:
    absent = fields - record.keys()
    if absent:
        raise ValueError(f"{sheet} is missing required columns: {', '.join(sorted(absent))}")


def load_data(workbook_path: Path) -> DataSet:
    """Load a workbook in read-only mode; source cells are never written by this workflow."""
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    expected = {"README", "Orders", "Shipments", "Inventory", "Operational_Updates"}
    if not expected.issubset(workbook.sheetnames):
        raise ValueError(f"Input workbook needs sheets: {', '.join(sorted(expected))}")
    settings = {}
    for row in workbook["README"].iter_rows(min_row=2, values_only=True):
        if row and row[0] and row[1] is not None:
            settings[str(row[0]).strip()] = row[1]
    as_of = _as_date(settings.get("planning_as_of"), "planning_as_of")
    actuals_through = _as_date(settings.get("actuals_through"), "actuals_through")
    if actuals_through < as_of:
        raise ValueError("actuals_through cannot precede planning_as_of")

    inventory_records = _records(workbook["Inventory"])
    required_item = {"item_id", "description", "category", "current_on_hand", "allocated_quantity",
                     "minimum_stock", "target_stock", "supplier_lead_time_days", "unit_cost",
                     "criticality", "last_updated"}
    items = tuple(Item(
        item_id=str(r["item_id"]).strip(), description=str(r["description"]).strip(),
        category=str(r["category"]).strip(), current_on_hand=number(r["current_on_hand"]),
        allocated_quantity=number(r["allocated_quantity"]), minimum_stock=number(r["minimum_stock"]),
        target_stock=number(r["target_stock"]),
        supplier_lead_time_days=int(r["supplier_lead_time_days"]) if r["supplier_lead_time_days"] not in (None, "") else None,
        unit_cost=number(r["unit_cost"]) if r["unit_cost"] not in (None, "") else None,
        criticality=str(r["criticality"]).strip(), last_updated=_as_date(r["last_updated"], "last_updated"),
    ) for r in inventory_records if not _required(r, required_item, "Inventory"))

    order_records = _records(workbook["Orders"])
    required_order = {"order_id", "order_date", "customer", "category", "item_id", "quantity", "unit_price", "status", "requested_ship_date"}
    orders = tuple(Order(
        order_id=str(r["order_id"]).strip(), order_date=_as_date(r["order_date"], "order_date"),
        customer=str(r["customer"]).strip(), category=str(r["category"]).strip(), item_id=str(r["item_id"]).strip(),
        quantity=number(r["quantity"]), unit_price=number(r["unit_price"]), status=str(r["status"]).strip().lower(),
        requested_ship_date=_as_date(r["requested_ship_date"], "requested_ship_date", optional=True),
    ) for r in order_records if not _required(r, required_order, "Orders"))

    shipment_records = _records(workbook["Shipments"])
    required_shipment = {"shipment_id", "order_id", "item_id", "ship_date", "quantity"}
    shipments = tuple(Shipment(
        shipment_id=str(r["shipment_id"]).strip(), order_id=str(r["order_id"]).strip(), item_id=str(r["item_id"]).strip(),
        ship_date=_as_date(r["ship_date"], "ship_date"), quantity=number(r["quantity"]),
    ) for r in shipment_records if not _required(r, required_shipment, "Shipments"))

    po_records = _records(workbook["Operational_Updates"])
    required_po = {"update_type", "reference", "item_id", "quantity", "expected_receipt_date", "status", "note"}
    purchase_orders = tuple(PurchaseOrder(
        str(r["reference"]).strip(), str(r["item_id"]).strip(), number(r["quantity"]),
        _as_date(r["expected_receipt_date"], "expected_receipt_date", optional=True), str(r["status"]).strip().lower(),
    ) for r in po_records if not _required(r, required_po, "Operational_Updates") and str(r["update_type"]).strip() == "purchase_order")
    workbook.close()

    usage_by_month: dict[tuple[str, date], Decimal] = defaultdict(Decimal)
    for order in orders:
        if order.order_date <= as_of and order.status != "cancelled":
            month = order.order_date.replace(day=1)
            usage_by_month[order.item_id, month] += order.quantity
    usage = tuple(Usage(item_id, month, quantity)
                  for (item_id, month), quantity in sorted(usage_by_month.items()))
    data = DataSet(items, usage, purchase_orders, orders, shipments, actuals_through, as_of)
    check_records(data)
    return data


def check_records(data: DataSet) -> None:
    for label, keys in (
        ("item", [x.item_id for x in data.items]),
        ("usage month", [(x.item_id, x.month) for x in data.usage]),
        ("PO reference", [x.po_reference for x in data.purchase_orders]),
        ("order", [x.order_id for x in data.orders]),
        ("shipment", [x.shipment_id for x in data.shipments]),
    ):
        if len(keys) != len(set(keys)):
            raise ValueError(f"Duplicate {label}; resolve before planning")
    item_ids = {x.item_id for x in data.items}
    if any(x.item_id not in item_ids for x in (*data.usage, *data.purchase_orders, *data.orders, *data.shipments)):
        raise ValueError("Usage, order, shipment, or PO references an unknown item")
    if any(x.month.day != 1 for x in data.usage):
        raise ValueError("Usage month must be its first calendar day")
    if any(x.status not in {"confirmed", "unconfirmed", "received", "cancelled"} for x in data.purchase_orders):
        raise ValueError("Unknown PO status; resolve before planning")
    if any(o.status not in {"open", "shipped", "complete", "cancelled"} for o in data.orders):
        raise ValueError("Unknown order status; resolve before analysis")
    if any(o.quantity < 0 or o.unit_price < 0 for o in data.orders) or any(s.quantity < 0 for s in data.shipments):
        raise ValueError("Order and shipment quantities/prices cannot be negative")
    order_ids = {o.order_id for o in data.orders}
    if any(s.order_id not in order_ids for s in data.shipments):
        raise ValueError("Shipment references an unknown order")
