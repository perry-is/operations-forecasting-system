"""Offline scenario using exclusively fictional records."""

import argparse
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from .excel import SHEET_ORDER, generate_workbook
from .ingest import load_data
from .reporting import summary, write_outputs
from .workflow import plan

ROOT = Path(__file__).resolve().parents[2]
AS_OF = date(2026, 6, 30)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-workbook", type=Path, default=ROOT / "examples" / "logistics_operations_input.xlsx")
    parser.add_argument("--output-dir", type=Path, default=ROOT / "generated_example")
    parser.add_argument("--history-file", type=Path)
    parser.add_argument("--horizon-days", type=int, default=90)
    parser.add_argument("--safety-factor", type=Decimal, default=Decimal("1.25"))
    parser.add_argument("--forecast-periods", type=int, default=3)
    args = parser.parse_args()
    if args.forecast_periods < 1:
        parser.error("--forecast-periods must be positive")
    source = args.input_workbook.resolve()
    data = load_data(source)
    as_of = data.planning_as_of or AS_OF
    result = plan(data, as_of,
                  horizon_days=args.horizon_days, safety_factor=args.safety_factor)
    write_outputs(result, args.output_dir)
    history_file = args.history_file or args.output_dir / "forecast_history.jsonl"
    workbook_path = args.output_dir / "logistics_forecast.xlsx"
    generated = generate_workbook(data, result, history_file, workbook_path, args.forecast_periods)
    reopened = load_workbook(workbook_path, read_only=True, data_only=True)
    if reopened.sheetnames != SHEET_ORDER or any(reopened[name].max_row < 2 for name in SHEET_ORDER):
        reopened.close()
        raise RuntimeError("Generated forecast workbook failed the output-sheet validation")
    output_sheets = list(reopened.sheetnames)
    reopened.close()
    baseline = [r for r in result["recommendations"] if r["scenario"] == "BASELINE"]
    print("Synthetic two-workbook workflow complete: input read-only; no purchasing actions executed.")
    print(f"Input workbook: {source} ({', '.join(generated['input_sheets'])})")
    print(f"Generated workbook: {workbook_path} ({', '.join(output_sheets)})")
    for key, value in summary(baseline).items():
        print(f"{key}: {value}")
    print(f"Forecast history: {len(generated['history'])} preserved records; {len(generated['history_conflicts'])} changed-ID candidates retained as review conflicts")
    print(f"Forecast-vs-actual: {sum(x['actual_quantity'] is not None for x in generated['forecast_vs_actual'])} periods scored")
    print("Outputs: logistics_forecast.xlsx, forecast_history.jsonl, recommendations.json, planning_report.md, audit.jsonl")


if __name__ == "__main__":
    main()
