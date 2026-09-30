"""Write the generated analysis workbook without touching the source workbook."""

import json
import math
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .analysis import historical_analysis
from .forecast_history import compare_forecasts_to_actuals, monthly_forecasts, preserve_forecast_history
from .models import DataSet
from .reporting import json_default, summary

SHEET_ORDER = ["Dashboard", "README", "Clean_Data", "Weekly_Analysis", "Monthly_Analysis",
               "Customer_Category_Analysis", "Inventory_Demand", "Forecast", "Forecast_History",
               "Forecast_vs_Actual", "Scenario_Analysis", "Recommendations", "Review_Queue", "Audit"]


def _plain(value):
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, (date, str, int, float, bool)) or value is None:
        return value
    return json.dumps(value, default=json_default, sort_keys=True)


def _column_value(key: str, value):
    if isinstance(value, str) and key in {"month", "week_start", "target_month", "forecast_origin", "as_of"}:
        try:
            return date.fromisoformat(value)
        except ValueError:
            pass
    return _plain(value)


def _table(workbook: Workbook, name: str, records: list[dict], columns: list[str] | None = None) -> None:
    sheet = workbook.create_sheet(name)
    if columns is None:
        columns = list(records[0]) if records else ["status"]
    sheet.append(columns)
    for record in records:
        sheet.append([_column_value(key, record.get(key)) for key in columns])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    header_fill = PatternFill("solid", fgColor="27445D")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10, color="23313D")
            cell.alignment = Alignment(vertical="center", wrap_text=isinstance(cell.value, str))
            if isinstance(cell.value, date):
                cell.number_format = "mm/dd/yy"
            elif isinstance(cell.value, float):
                header = sheet.cell(1, cell.column).value
                cell.number_format = '0.0"%"' if header == "percentage_error" else "#,##0.00"
    for index, column in enumerate(sheet.columns, start=1):
        lengths = [len(str(cell.value)) if cell.value is not None else 0 for cell in column[:100]]
        sheet.column_dimensions[get_column_letter(index)].width = min(max(max(lengths, default=8) + 2, 12), 36)
    for row_index, row in enumerate(sheet.iter_rows(min_row=2), start=2):
        line_count = max((math.ceil(len(str(cell.value)) / max(sheet.column_dimensions[get_column_letter(cell.column)].width - 2, 1))
                          for cell in row if isinstance(cell.value, str)), default=1)
        sheet.row_dimensions[row_index].height = min(max(18, line_count * 14), 56)
    sheet.row_dimensions[1].height = 30


def _recommendation_rows(result: dict) -> list[dict]:
    rows = []
    for record in result["recommendations"]:
        if record["scenario"] != "BASELINE":
            continue
        projection = record["projection"] or {}
        rec = record["recommendation"]
        rows.append({
            "item_id": record["item_id"], "description": record["description"], "criticality": rec["priority"],
            "current_on_hand": record["observed_facts"]["current_on_hand"],
            "current_available": projection.get("current_available"),
            "projected_at_lead_time": projection.get("projected_stock_at_lead_time"),
            "stockout_date": projection.get("stockout_date"),
            "action": rec["action"], "suggested_quantity": rec["suggested_quantity"],
            "reason": rec["reason"], "review_flags": rec["review_flags"],
            "review_required": rec["review_required"],
        })
    return rows


def generate_workbook(data: DataSet, result: dict, history_path: Path, output_path: Path,
                      forecast_periods: int = 3) -> dict:
    baseline = [r for r in result["recommendations"] if r["scenario"] == "BASELINE"]
    candidates = monthly_forecasts(baseline, result["as_of"], forecast_periods)
    history, collisions = preserve_forecast_history(history_path, candidates)
    comparisons = compare_forecasts_to_actuals(history, data.orders, data.actuals_through)
    analysis = historical_analysis(data, result["as_of"])

    orders = [o for o in data.orders if o.order_date <= (data.actuals_through or result["as_of"]) and o.status != "cancelled"]
    clean = [{"record_type": "order", "record_id": o.order_id, "record_date": o.order_date,
              "item_id": o.item_id, "customer": o.customer, "category": o.category,
              "quantity": o.quantity, "unit_price": o.unit_price, "extended_value": o.quantity * o.unit_price,
              "status": o.status} for o in orders]
    for shipment in data.shipments:
        if shipment.ship_date <= (data.actuals_through or result["as_of"]):
            clean.append({"record_type": "shipment", "record_id": shipment.shipment_id,
                          "record_date": shipment.ship_date, "item_id": shipment.item_id,
                          "customer": None, "category": None, "quantity": shipment.quantity,
                          "unit_price": None, "extended_value": None, "status": "shipped"})

    scenario_rows = []
    by_key = {(r["scenario"], r["item_id"]): r for r in result["recommendations"]}
    for baseline_row in baseline:
        item_id = baseline_row["item_id"]
        row = {"item_id": item_id}
        for scenario_name in result["scenario_names"]:
            scenario = by_key[scenario_name, item_id]
            rec = scenario["recommendation"]
            row[scenario_name] = f"{rec['action']}: {rec['suggested_quantity']}" if rec["suggested_quantity"] else rec["action"]
        scenario_rows.append(row)
    recommendation_rows = _recommendation_rows(result)
    review_rows = [r for r in recommendation_rows if r["review_required"] or r["review_flags"] or r["action"] == "HUMAN REVIEW"]
    candidate_by_id = {row["forecast_id"]: row for row in candidates}
    saved_by_id = {row["forecast_id"]: row for row in history}
    for forecast_id in collisions:
        saved, candidate = saved_by_id[forecast_id], candidate_by_id[forecast_id]
        review_rows.append({"item_id": saved["item_id"], "description": None, "criticality": None,
                            "current_on_hand": None, "current_available": None,
                            "projected_at_lead_time": None, "stockout_date": None,
                            "action": "HUMAN REVIEW", "suggested_quantity": None,
                            "reason": f"Forecast {forecast_id} changed from {saved['forecast_quantity']} to {candidate['forecast_quantity']}; the saved value remains authoritative.",
                            "review_flags": ["FORECAST_SNAPSHOT_CONFLICT"], "review_required": True})
    audit_rows = []
    for event in result["audit"]:
        if event["scenario"]["name"] != "BASELINE":
            continue
        item = event["source_values"]["item"]
        forecast = event["forecast"]
        projection = event["projected_inventory"] or {}
        rec = event["recommendation"]
        audit_rows.append({"as_of": event["as_of"], "scenario": event["scenario"]["name"],
                           "item_id": event["item_id"], "on_hand": item["current_on_hand"],
                           "allocated": item["allocated_quantity"], "usage_months": len(event["source_values"]["usage"]),
                           "purchase_orders": len(event["source_values"]["purchase_orders"]),
                           "forecast_method": forecast["forecast_method"],
                           "forecast_daily_usage": forecast["forecast_daily_usage"],
                           "forecast_period_usage": forecast["forecast_period_usage"],
                           "projected_at_lead_time": projection.get("projected_stock_at_lead_time"),
                           "stockout_date": projection.get("stockout_date"),
                           "action": rec["action"], "suggested_quantity": rec["suggested_quantity"],
                           "reason": rec["reason"], "review_flags": rec["review_flags"]})

    wb = Workbook()
    dash = wb.active
    dash.title = "Dashboard"
    baseline_summary = summary(baseline)
    dashboard_rows = [
        {"measure": "Planning snapshot", "value": result["as_of"], "interpretation": "Inventory view date"},
        {"measure": "Historical actuals through", "value": data.actuals_through, "interpretation": "Synthetic order and shipment records"},
        *[{"measure": key.replace("_", " ").title(), "value": value, "interpretation": "Baseline recommendation set"} for key, value in baseline_summary.items()],
        {"measure": "Forecast periods", "value": forecast_periods, "interpretation": "Monthly forecasts after the planning snapshot"},
        {"measure": "Preserved forecast records", "value": len(history), "interpretation": "First stored value retained for each forecast ID"},
        {"measure": "Forecast history conflicts", "value": len(collisions), "interpretation": "Same forecast ID produced a changed candidate; stored value kept"},
        {"measure": "Forecast comparisons with actuals", "value": sum(r["actual_quantity"] is not None for r in comparisons), "interpretation": "Complete months available through actuals cutoff"},
    ]
    _fill_existing(dash, ["measure", "value", "interpretation"], dashboard_rows)
    readme = [{"topic": "Purpose", "detail": "Generated analysis workbook from synthetic human-maintained input data."},
              {"topic": "Source", "detail": "Python reads examples/logistics_operations_input.xlsx in read-only mode. This workbook is never overwritten."},
              {"topic": "Planning snapshot", "detail": result["as_of"].isoformat()},
              {"topic": "Historical actual cutoff", "detail": data.actuals_through.isoformat() if data.actuals_through else "not supplied"},
              {"topic": "Forecast vs actual", "detail": "Absolute error is |forecast - actual|. Percentage error is unavailable when actual is zero."},
              {"topic": "Decision boundary", "detail": "Recommendations are evidence for human review. The workflow does not place orders."},
              {"topic": "Synthetic scope", "detail": "All items, customers, order references, shipments, dates, quantities, and costs are fictional."},
              {"topic": "Refresh", "detail": "Run operations-forecast-demo to read the input workbook and regenerate this analysis workbook."}]
    _table(wb, "README", readme)
    _table(wb, "Clean_Data", clean, ["record_type", "record_id", "record_date", "item_id", "customer", "category", "quantity", "unit_price", "extended_value", "status"])
    _table(wb, "Weekly_Analysis", analysis["weekly"])
    _table(wb, "Monthly_Analysis", analysis["monthly"])
    _table(wb, "Customer_Category_Analysis", analysis["customer_category"])
    _table(wb, "Inventory_Demand", analysis["inventory_demand"])
    # Show the value actually retained for this run's forecast IDs, not a conflicting recalculation.
    forecast_columns = ["forecast_id", "forecast_origin", "item_id", "target_month", "forecast_quantity",
                        "forecast_daily_usage", "forecast_method", "history_window_months", "confidence_note"]
    _table(wb, "Forecast", [saved_by_id[row["forecast_id"]] for row in candidates], forecast_columns)
    _table(wb, "Forecast_History", history, forecast_columns)
    _table(wb, "Forecast_vs_Actual", comparisons,
           forecast_columns + ["actual_quantity", "absolute_error", "percentage_error", "direction", "actual_available_through"])
    _table(wb, "Scenario_Analysis", scenario_rows)
    _table(wb, "Recommendations", recommendation_rows)
    _table(wb, "Review_Queue", review_rows)
    _table(wb, "Audit", audit_rows)
    # Keep the requested reader-facing order and a restrained consistent typeface.
    wb._sheets = [wb[name] for name in SHEET_ORDER]
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None and cell.font.name != "Arial":
                    cell.font = Font(name="Arial", size=10, bold=cell.font.bold, color=cell.font.color)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    wb.close()
    return {"analysis": analysis, "forecasts": candidates, "history": history,
            "forecast_vs_actual": comparisons, "history_conflicts": collisions,
            "input_sheets": ["README", "Orders", "Shipments", "Inventory", "Operational_Updates"],
            "output_sheets": SHEET_ORDER}


def _fill_existing(sheet, columns: list[str], records: list[dict]) -> None:
    for row in (columns, *[[ _plain(r.get(key)) for key in columns] for r in records]):
        sheet.append(list(row))
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="27445D")
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10, color="23313D")
            cell.alignment = Alignment(vertical="center", wrap_text=isinstance(cell.value, str))
            if isinstance(cell.value, date):
                cell.number_format = "mm/dd/yy"
    for index, column in enumerate(sheet.columns, start=1):
        lengths = [len(str(cell.value)) if cell.value is not None else 0 for cell in column[:100]]
        sheet.column_dimensions[get_column_letter(index)].width = min(max(max(lengths, default=8) + 2, 12), 38)
    for row_index, row in enumerate(sheet.iter_rows(min_row=2), start=2):
        line_count = max((math.ceil(len(str(cell.value)) / max(sheet.column_dimensions[get_column_letter(cell.column)].width - 2, 1))
                          for cell in row if isinstance(cell.value, str)), default=1)
        sheet.row_dimensions[row_index].height = min(max(18, line_count * 14), 56)
