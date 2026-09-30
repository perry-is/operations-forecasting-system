import copy
import json
import hashlib
import shutil
import subprocess
import sys
import tempfile
import unittest
from dataclasses import replace
from datetime import date
from decimal import Decimal as D
from pathlib import Path

from operations_forecasting.demo import AS_OF, ROOT
from operations_forecasting.forecasting import forecast
from operations_forecasting.analysis import historical_analysis
from operations_forecasting.excel import SHEET_ORDER
from operations_forecasting.forecast_history import compare_forecasts_to_actuals, monthly_forecasts, preserve_forecast_history
from operations_forecasting.ingest import check_records, load_data, number
from operations_forecasting.intake import (ApprovedOperationalRecord, ProposedOperationalRecord,
                                           admit_approved_records)
from operations_forecasting.models import DataSet, PurchaseOrder, Scenario, Usage
from operations_forecasting.reporting import render_report, summary, write_outputs
from operations_forecasting.workflow import plan
from openpyxl import load_workbook


class PlanningTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.input_workbook = ROOT / "examples" / "logistics_operations_input.xlsx"
        cls.data = load_data(cls.input_workbook)
        cls.result = plan(cls.data, AS_OF)
        cls.index = {(r["scenario"], r["item_id"]): r for r in cls.result["recommendations"]}

    def row(self, item_id="PART-001", scenario="BASELINE"):
        return self.index[scenario, item_id]

    def one(self, item, orders=(), history=None):
        if history is None:
            history = tuple(x for x in self.data.usage if x.item_id == item.item_id)
        return plan(DataSet((item,), history, orders), AS_OF, (Scenario(),))["recommendations"][0]

    def test_moving_average_uses_observed_calendar_days(self):
        f = self.row()["forecast"]
        self.assertEqual(f["observed_days"], 91)
        self.assertEqual(f["observed_usage"], D(91))
        self.assertEqual(f["forecast_daily_usage"], D(1))
        self.assertEqual(f["forecast_period_usage"], D(90))

    def test_inventory_position_and_allocations(self):
        p = self.row()["projection"]
        self.assertEqual(p["current_on_hand"], D(120))
        self.assertEqual(p["current_available"], D(110))
        self.assertEqual(p["projected_stock_at_lead_time"], D(96))
        self.assertEqual(p["projected_available_at_horizon"], D(20))

    def test_allocation_change_does_not_change_forecast(self):
        item = self.data.items[0]
        a, b = self.one(item), self.one(replace(item, allocated_quantity=D(20)))
        self.assertEqual(a["forecast"], b["forecast"])
        self.assertEqual(a["projection"]["current_available"] - b["projection"]["current_available"], D(10))

    def test_inbound_prevents_false_reorder_without_becoming_on_hand(self):
        row = self.row("PART-004")
        self.assertEqual(row["projection"]["current_on_hand"], D(25))
        self.assertEqual(row["projection"]["confirmed_inbound_by_lead_time"], D(60))
        self.assertEqual(row["projection"]["projected_stock_at_lead_time"], D(55))
        self.assertEqual(row["recommendation"]["action"], "MONITOR")
        item = next(x for x in self.data.items if x.item_id == "PART-004")
        self.assertEqual(self.one(item)["recommendation"]["action"], "EXPEDITE")

    def test_overdue_po_is_excluded_and_quantity_withheld(self):
        row = self.row("PART-005")
        self.assertEqual(row["projection"]["confirmed_inbound_within_horizon"], 0)
        self.assertIn("OVERDUE_PURCHASE_ORDER", row["recommendation"]["review_flags"])
        self.assertEqual(row["recommendation"]["action"], "HUMAN REVIEW")
        self.assertIsNone(row["recommendation"]["suggested_quantity"])

    def test_currently_healthy_stock_can_run_out_during_lead_time(self):
        row = self.row("PART-003")
        self.assertEqual(row["projection"]["current_available"], D(100))
        self.assertEqual(row["projection"]["projected_stock_at_lead_time"], D(-20))
        self.assertEqual(row["recommendation"]["action"], "EXPEDITE")
        self.assertEqual(row["recommendation"]["suggested_quantity"], 200)
        self.assertTrue(row["recommendation"]["lead_time_risk"])

    def test_reorder_trigger_and_order_quantity_are_explainable(self):
        row = self.row("PART-002")
        self.assertEqual(row["forecast"]["forecast_daily_usage"], D(2))
        self.assertEqual(row["projection"]["projected_stock_at_lead_time"], D(20))
        self.assertEqual(row["recommendation"]["action"], "ORDER")
        self.assertEqual(row["recommendation"]["suggested_quantity"], 90)
        self.assertEqual(row["recommendation"]["quantity_calculation"]["target_gap"], D(90))

    def test_stable_healthy_item_needs_no_action(self):
        self.assertEqual(self.row()["recommendation"]["action"], "NO ACTION")
        self.assertIsNone(self.row()["projection"]["stockout_date"])

    def test_stockout_is_strictly_below_zero(self):
        self.assertEqual(self.row("PART-003")["projection"]["stockout_date"], date(2026, 8, 3))
        self.assertEqual(self.row("PART-011")["projection"]["stockout_date"], date(2026, 8, 22))

    def test_risk_threshold_dates_are_distinct(self):
        p = self.row("PART-011")["projection"]
        self.assertEqual(p["below_safety_date"], date(2026, 8, 9))
        self.assertEqual(p["below_minimum_date"], date(2026, 8, 12))
        self.assertLess(p["below_minimum_date"], p["stockout_date"])

    def test_sparse_history_is_not_padded_with_zero_months(self):
        row = self.row("PART-009")
        self.assertEqual(row["forecast"]["observed_days"], 30)
        self.assertEqual(row["forecast"]["forecast_daily_usage"], D(1))
        self.assertIn("INSUFFICIENT_HISTORY", row["recommendation"]["review_flags"])
        self.assertIsNone(row["recommendation"]["suggested_quantity"])

    def test_no_history_has_no_fabricated_forecast_or_projection(self):
        row = self.one(self.data.items[0], history=())
        self.assertIsNone(row["forecast"]["forecast_daily_usage"])
        self.assertIsNone(row["projection"])
        self.assertEqual(row["recommendation"]["action"], "HUMAN REVIEW")
        self.assertIsNone(row["recommendation"]["lead_time_risk"])
        self.assertIsNone(row["recommendation"]["excess_stock"])

    def test_report_keeps_known_availability_and_marks_unavailable_risk_unknown(self):
        row = self.one(self.data.items[0], history=())
        report = render_report({"as_of": AS_OF, "scenario_names": ["BASELINE"], "recommendations": [row]})
        self.assertIn("| PART-001 | 2 | 110 | unknown | unknown | HUMAN REVIEW | withheld", report)
        self.assertNotIn("none in horizon", report)

    def test_missing_lead_time_and_stale_count_require_review(self):
        row = self.row("KIT-002")
        flags = row["recommendation"]["review_flags"]
        for expected in ("UNKNOWN_LEAD_TIME", "STALE_COUNT", "INCONSISTENT_QUANTITIES", "MISSING_RECEIPT_DATE"):
            self.assertIn(expected, flags)
        self.assertIsNone(row["projection"]["projected_stock_at_lead_time"])
        self.assertIsNone(row["recommendation"]["candidate_quantity"])
        self.assertIsNone(row["recommendation"]["lead_time_risk"])

    def test_negative_inventory_requires_review(self):
        row = self.one(replace(self.data.items[0], current_on_hand=D(-2), allocated_quantity=D(0)))
        self.assertIn("NEGATIVE_INVENTORY", row["recommendation"]["review_flags"])
        self.assertEqual(row["projection"]["stockout_date"], AS_OF)

    def test_spike_and_intermittent_usage_are_not_confident(self):
        self.assertIn("DEMAND_ANOMALY", self.row("PART-007")["recommendation"]["review_flags"])
        self.assertIn("INTERMITTENT_DEMAND", self.row("PART-006")["recommendation"]["review_flags"])
        self.assertEqual(self.row("PART-007")["forecast"]["data_quality"], "low")

    def test_excess_and_zero_recent_demand_are_reviewed_without_disposal(self):
        for item_id in ("PART-008", "KIT-001"):
            self.assertEqual(self.row(item_id)["recommendation"]["action"], "REDUCE / REVIEW EXCESS")
            self.assertTrue(self.row(item_id)["recommendation"]["review_required"])
        self.assertEqual(self.row("KIT-001")["forecast"]["forecast_daily_usage"], D(0))

    def test_scenarios_do_not_mutate_source_records(self):
        original = copy.deepcopy(self.data)
        self.assertEqual(plan(self.data, AS_OF), self.result)
        self.assertEqual(self.data, original)
        for row in self.result["recommendations"]:
            self.assertEqual(row["observed_facts"], self.row(row["item_id"])["observed_facts"])

    def test_demand_increase_changes_projection_and_recommendation(self):
        base, changed = self.row("PART-011"), self.row("PART-011", "DEMAND +20%")
        self.assertEqual(changed["forecast"]["forecast_daily_usage"], D("1.2"))
        self.assertEqual(changed["projection"]["projected_stock_at_lead_time"], D(16))
        self.assertEqual(base["recommendation"]["action"], "MONITOR")
        self.assertEqual(changed["recommendation"]["action"], "ORDER")
        self.assertEqual(changed["recommendation"]["suggested_quantity"], 44)

    def test_lead_time_delay_changes_order_timing_and_quantity(self):
        base, changed = self.row("PART-012"), self.row("PART-012", "LEAD TIME +14 DAYS")
        self.assertEqual(base["recommendation"]["action"], "MONITOR")
        self.assertEqual(changed["recommendation"]["action"], "ORDER")
        self.assertEqual(changed["recommendation"]["lead_time_days"], 49)
        self.assertEqual(changed["recommendation"]["suggested_quantity"], 59)
        self.assertEqual(base["forecast"]["forecast_daily_usage"], changed["forecast"]["forecast_daily_usage"])

    def test_delayed_po_changes_risk_but_does_not_rescue_overdue_po(self):
        delayed = self.row("PART-004", "INBOUND PO DELAYED")
        self.assertEqual(delayed["recommendation"]["action"], "EXPEDITE")
        self.assertEqual(delayed["projection"]["stockout_date"], date(2026, 7, 26))
        self.assertEqual(self.row("PART-005", "INBOUND PO DELAYED")["projection"]["confirmed_inbound_within_horizon"], 0)

    def test_criticality_changes_priority_not_forecast(self):
        item = self.data.items[0]
        low, critical = self.one(replace(item, criticality="low")), self.one(replace(item, criticality="critical"))
        self.assertEqual(low["forecast"], critical["forecast"])
        self.assertEqual(low["projection"], critical["projection"])
        self.assertLess(low["recommendation"]["priority"], critical["recommendation"]["priority"])

    def test_audit_contains_inputs_assumptions_and_explanation(self):
        event = self.result["audit"][0]
        for key in ("item_id", "source_values", "scenario", "forecast", "projected_inventory", "inbound_supply_considered", "recommendation", "policy"):
            self.assertIn(key, event)
        self.assertEqual(event["recommendation"]["human_decision"]["status"], "not_recorded")

    def test_duplicate_po_is_rejected_instead_of_double_counted(self):
        with self.assertRaisesRegex(ValueError, "Duplicate PO reference"):
            check_records(replace(self.data, purchase_orders=self.data.purchase_orders + self.data.purchase_orders[:1]))

    def test_received_cancelled_unconfirmed_orders_are_excluded(self):
        self.assertEqual(self.row()["projection"]["confirmed_inbound_within_horizon"], 0)
        self.assertEqual(self.row("PART-004")["projection"]["confirmed_inbound_within_horizon"], 60)
        self.assertEqual(self.row("PART-008")["projection"]["confirmed_inbound_within_horizon"], 0)

    def test_receipts_after_horizon_are_not_counted_early(self):
        item = self.data.items[0]
        po = PurchaseOrder("PO-SYN-TEST", item.item_id, D(1000), date(2026, 12, 1), "confirmed")
        row = self.one(item, (po,))
        self.assertEqual(row["projection"]["projected_available_at_horizon"], D(20))

    def test_receipt_on_stockout_day_is_available_before_daily_usage(self):
        item = replace(self.data.items[0], current_on_hand=D(10), allocated_quantity=D(0))
        po = PurchaseOrder("PO-SYN-TEST", item.item_id, D(100), date(2026, 7, 11), "confirmed")
        self.assertIsNone(self.one(item, (po,))["projection"]["stockout_date"])

    def test_inbound_recovery_never_creates_zero_quantity_order(self):
        item = replace(self.data.items[0], current_on_hand=D(12), allocated_quantity=D(0))
        po = PurchaseOrder("PO-SYN-TEST", item.item_id, D(100), date(2026, 7, 2), "confirmed")
        self.assertEqual(self.one(item, (po,))["recommendation"]["action"], "MONITOR")

    def test_invalid_numeric_input_and_scenario_are_rejected(self):
        for value in ("NaN", "Infinity"):
            with self.assertRaises(ValueError):
                number(value)
        with self.assertRaises(ValueError):
            Scenario(demand_multiplier=D(-1))

    def test_incomplete_current_month_is_not_treated_as_complete_history(self):
        history = tuple(x for x in self.data.usage if x.item_id == "PART-001")
        value = forecast(history, date(2026, 6, 15))
        self.assertEqual(value["observed_months"], [date(2026, 3, 1), date(2026, 4, 1), date(2026, 5, 1)])
        self.assertEqual(value["baseline_daily_usage"], D(1))

    def test_summary_excludes_withheld_candidates_from_purchase_value(self):
        baseline = [r for r in self.result["recommendations"] if r["scenario"] == "BASELINE"]
        self.assertEqual(summary(baseline)["suggested_purchase_value"], D("2517.00"))

    def test_end_to_end_cli_produces_expected_outputs_and_scenarios(self):
        with tempfile.TemporaryDirectory() as temp:
            completed = subprocess.run([sys.executable, "-m", "operations_forecasting.demo", "--output-dir", temp], capture_output=True, text=True, timeout=30)
            self.assertEqual(completed.returncode, 0, completed.stderr)
            folder = Path(temp)
            self.assertEqual({p.name for p in folder.iterdir()}, {"recommendations.json", "planning_report.md", "audit.jsonl",
                                                                  "logistics_forecast.xlsx", "forecast_history.jsonl"})
            output = json.loads((folder / "recommendations.json").read_text(encoding="utf-8"))
            self.assertEqual(len(output["recommendations"]), 56)
            self.assertEqual(len(output["scenario_names"]), 4)
            events = [json.loads(line) for line in (folder / "audit.jsonl").read_text(encoding="utf-8").splitlines()]
            self.assertEqual(len(events), 56)
            before = {p.name: p.read_bytes() for p in folder.iterdir()}
            write_outputs(self.result, folder)
            self.assertEqual(before, {p.name: p.read_bytes() for p in folder.iterdir()})


class WorkbookWorkflowTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.source = ROOT / "examples" / "logistics_operations_input.xlsx"
        cls.data = load_data(cls.source)
        cls.result = plan(cls.data, AS_OF)

    def run_demo(self, directory: str, source: Path | None = None):
        source = source or self.source
        return subprocess.run([sys.executable, "-m", "operations_forecasting.demo", "--input-workbook", str(source),
                               "--output-dir", directory], capture_output=True, text=True, timeout=30)

    def test_input_workbook_structure(self):
        wb = load_workbook(self.source, read_only=True, data_only=True)
        self.assertEqual(wb.sheetnames, ["README", "Orders", "Shipments", "Inventory", "Operational_Updates"])
        self.assertEqual(sum(1 for _ in wb["Orders"].iter_rows()), 243)
        self.assertEqual(sum(1 for _ in wb["Inventory"].iter_rows()), 15)
        wb.close()

    def test_workbook_order_and_shipment_data_is_used(self):
        self.assertEqual(len(self.data.orders), 242)
        self.assertEqual(len(self.data.shipments), 242)
        self.assertEqual(len(self.data.usage), 79)
        june = next(x for x in self.data.usage if x.item_id == "PART-002" and x.month == date(2026, 6, 1))
        self.assertEqual(june.quantity, D(75))

    def test_weekly_analysis_has_order_and_shipping_progression(self):
        weekly = historical_analysis(self.data, AS_OF)["weekly"]
        self.assertGreater(len(weekly), 10)
        self.assertTrue(any(row["ordered_quantity"] > 0 and row["shipped_quantity"] > 0 for row in weekly))
        self.assertEqual(sum((row["ordered_quantity"] for row in weekly), D(0)),
                         sum((o.quantity for o in self.data.orders), D(0)))

    def test_monthly_analysis_and_demand_progression(self):
        analysis = historical_analysis(self.data, AS_OF)
        months = {row["month"] for row in analysis["monthly"]}
        self.assertEqual(months, {date(2026, m, 1) for m in range(1, 10)})
        p002 = [r["ordered_quantity"] for r in analysis["inventory_demand"] if r["item_id"] == "PART-002"]
        self.assertEqual(p002[:6], [D(15), D(20), D(31), D(45), D(62), D(75)])
        self.assertGreater(p002[-1], p002[0])

    def test_customer_and_category_patterns_are_grouped(self):
        rows = historical_analysis(self.data, AS_OF)["customer_category"]
        self.assertGreaterEqual(len(rows), 4)
        self.assertEqual({r["customer"] for r in rows}, {"Fictional Customer North", "Fictional Customer South"})

    def test_forecast_history_preserves_original_values(self):
        baseline = [r for r in self.result["recommendations"] if r["scenario"] == "BASELINE"]
        candidates = monthly_forecasts(baseline, AS_OF)
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "history.jsonl"
            original, collisions = preserve_forecast_history(path, candidates)
            altered = [{**candidates[0], "forecast_quantity": "99999"}]
            second, collisions = preserve_forecast_history(path, altered)
            self.assertEqual(second[0]["forecast_quantity"], original[0]["forecast_quantity"])
            self.assertIn(candidates[0]["forecast_id"], collisions)
            self.assertEqual(len(second), len(original))

    def test_history_records_are_not_rewritten_when_source_candidate_changes(self):
        baseline = [r for r in self.result["recommendations"] if r["scenario"] == "BASELINE"]
        forecasts = monthly_forecasts(baseline, AS_OF)
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "history.jsonl"
            preserve_forecast_history(path, forecasts)
            first_bytes = path.read_bytes()
            changed = [{**forecasts[0], "forecast_daily_usage": "999", "forecast_quantity": "99999"}]
            preserve_forecast_history(path, changed)
            second_bytes = path.read_bytes()
            self.assertEqual(first_bytes, second_bytes)
            saved = next(json.loads(line) for line in path.read_text(encoding="utf-8").splitlines()
                         if json.loads(line)["forecast_id"] == forecasts[0]["forecast_id"])
            self.assertEqual(saved["forecast_quantity"], forecasts[0]["forecast_quantity"])

    def test_later_actual_update_does_not_rewrite_saved_forecast(self):
        with tempfile.TemporaryDirectory() as temp:
            temp_path = Path(temp)
            source_copy = temp_path / "input.xlsx"
            shutil.copy2(self.source, source_copy)
            first_run = self.run_demo(str(temp_path / "out"), source_copy)
            self.assertEqual(first_run.returncode, 0, first_run.stderr)
            history_file = temp_path / "out" / "forecast_history.jsonl"
            saved_before = {r["forecast_id"]: r["forecast_quantity"] for r in
                            (json.loads(line) for line in history_file.read_text(encoding="utf-8").splitlines())}
            wb = load_workbook(source_copy)
            sheet = wb["Orders"]
            headers = {cell.value: cell.column for cell in sheet[1]}
            for row in range(2, sheet.max_row + 1):
                if str(sheet.cell(row, headers["item_id"]).value) == "PART-001" and sheet.cell(row, headers["order_date"]).value.month == 7:
                    sheet.cell(row, headers["quantity"]).value *= 2
            wb.save(source_copy)
            wb.close()
            second_run = self.run_demo(str(temp_path / "out"), source_copy)
            self.assertEqual(second_run.returncode, 0, second_run.stderr)
            saved_after = {r["forecast_id"]: r["forecast_quantity"] for r in
                           (json.loads(line) for line in history_file.read_text(encoding="utf-8").splitlines())}
            self.assertEqual(saved_before, saved_after)
            output = load_workbook(temp_path / "out" / "logistics_forecast.xlsx", read_only=True, data_only=True)
            rows = list(output["Forecast_vs_Actual"].iter_rows(values_only=True))
            output.close()
            columns = {name: index for index, name in enumerate(rows[0])}
            p001_july = next(row for row in rows[1:] if row[columns["item_id"]] == "PART-001" and
                             (row[columns["target_month"]].date() if hasattr(row[columns["target_month"]], "date") else row[columns["target_month"]]) == date(2026, 7, 1))
            self.assertEqual(p001_july[columns["actual_quantity"]], 60)

    def test_forecast_vs_actual_includes_errors_and_direction(self):
        baseline = [r for r in self.result["recommendations"] if r["scenario"] == "BASELINE"]
        forecasts = monthly_forecasts(baseline, AS_OF)
        comparisons = compare_forecasts_to_actuals(forecasts, self.data.orders, self.data.actuals_through)
        july = next(r for r in comparisons if r["item_id"] == "PART-001" and r["target_month"] == "2026-07-01")
        self.assertEqual(july["actual_quantity"], D(30))
        self.assertEqual(july["absolute_error"], abs(D(july["forecast_quantity"]) - D(30)))
        self.assertIn(july["direction"], {"over", "under", "even"})
        self.assertIsNotNone(july["percentage_error"])

    def test_forecast_vs_actual_percentage_error_is_undefined_when_actual_is_zero(self):
        forecast = {"forecast_id":"x", "forecast_origin":"2026-06-30", "item_id":"KIT-001",
                    "target_month":"2026-07-01", "forecast_quantity":"0", "forecast_method":"test",
                    "history_window_months":3, "forecast_daily_usage":"0", "confidence_note":"synthetic"}
        orders = tuple(o for o in self.data.orders if o.item_id == "KIT-001" and o.order_date.month == 7)
        compared = compare_forecasts_to_actuals([forecast], orders, date(2026, 9, 30))[0]
        self.assertEqual(compared["actual_quantity"], D(0))
        self.assertIsNone(compared["percentage_error"])

    def test_proposed_intake_cannot_enter_canonical_data(self):
        proposal = ProposedOperationalRecord("proposal-1", "fictional-doc-1", (("quantity", "4"),),
                                             ("quantity: 4",), ("unit unclear",))
        with self.assertRaisesRegex(TypeError, "Only human-approved"):
            admit_approved_records([proposal])

    def test_human_approved_or_corrected_intake_can_proceed(self):
        approved = ApprovedOperationalRecord("fact-1", (("quantity", "4"),), "corrected", ("fictional-doc-1",))
        self.assertEqual(admit_approved_records([approved]), (approved,))

    def test_demo_leaves_human_source_workbook_unchanged(self):
        before = hashlib.sha256(self.source.read_bytes()).hexdigest()
        with tempfile.TemporaryDirectory() as temp:
            completed = self.run_demo(temp)
            self.assertEqual(completed.returncode, 0, completed.stderr)
        after = hashlib.sha256(self.source.read_bytes()).hexdigest()
        self.assertEqual(before, after)

    def test_generated_forecast_workbook_has_populated_sheets_and_reopens(self):
        with tempfile.TemporaryDirectory() as temp:
            completed = self.run_demo(temp)
            self.assertEqual(completed.returncode, 0, completed.stderr)
            path = Path(temp) / "logistics_forecast.xlsx"
            wb = load_workbook(path, read_only=True, data_only=True)
            self.assertEqual(wb.sheetnames, SHEET_ORDER)
            for name in SHEET_ORDER:
                self.assertGreaterEqual(wb[name].max_row, 2, name)
            self.assertGreater(wb["Weekly_Analysis"].max_row, 10)
            self.assertGreater(wb["Forecast_vs_Actual"].max_row, 2)
            wb.close()

    def test_complete_demo_reads_workbook_and_generates_second_workbook(self):
        with tempfile.TemporaryDirectory() as temp:
            completed = self.run_demo(temp)
            self.assertEqual(completed.returncode, 0, completed.stderr)
            self.assertIn("Input workbook:", completed.stdout)
            self.assertIn("Generated workbook:", completed.stdout)
            self.assertIn("42 preserved records", completed.stdout)
            self.assertTrue((Path(temp) / "logistics_forecast.xlsx").exists())
            self.assertTrue((Path(temp) / "forecast_history.jsonl").exists())


if __name__ == "__main__":
    unittest.main()
